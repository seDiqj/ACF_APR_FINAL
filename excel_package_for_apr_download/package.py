from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# -------------------------
# داده نمونه
# -------------------------
final_apr = {
    "impact": "Vulnerable receiving focused MHPSS care",
    "outcomes": [
        {
            "name": "Communities have increased access to MHPSS services",
            "outputs": [
                {
                    "name": "Output 1: Hotline & remote services",
                    "indicators": [
                        {
                            "code": "1.1.1",
                            "name": "Dedicated hotline - psychosocial support",
                            "disaggregation": [
                                {"name": "# of Male (above 18)", "target": 30, "months": [5, 8, 13, 2, 3, 0, 1, 0, 0, 2, 1, 1]},
                                {"name": "# of Female (above 18)", "target": 1200, "months": [400, 390, 398, 5, 5, 0, 1, 1, 0, 1, 1, 0]},
                                {"name": "# of Male under 18", "target": 40, "months": [5, 8, 8, 2, 2, 0, 0, 0, 0, 1, 1, 0]},
                            ]
                        },
                        {
                            "code": "1.1.2",
                            "name": "Women friendly spaces - MHPSS sessions",
                            "disaggregation": [
                                {"name": "# of Female (above 18)", "target": 1458, "months": [450, 500, 508, 0, 0, 0, 0, 0, 0, 0, 0, 0]},
                                {"name": "# of Male (above 18)", "target": 32, "months": [10, 12, 10, 0, 0, 0, 0, 0, 0, 0, 0, 0]},
                                {"name": "# of Female under 18", "target": 50, "months": [5, 6, 4, 0, 0, 0, 0, 0, 0, 0, 0, 0]},
                                {"name": "# of Male under 18", "target": 20, "months": [2, 2, 3, 0, 0, 0, 0, 0, 0, 0, 0, 0]},
                            ]
                        },
                    ]
                },
                {
                    "name": "Output 2: Community outreach",
                    "indicators": [
                        {
                            "code": "1.2.1",
                            "name": "Awareness sessions in schools",
                            "disaggregation": [
                                {"name": "Boys", "target": 120, "months": [40, 40, 40, 0, 0, 0, 0, 0, 0, 0, 0, 0]},
                                {"name": "Girls", "target": 130, "months": [40, 45, 45, 0, 0, 0, 0, 0, 0, 0, 0, 0]},
                            ]
                        },
                        {
                            "code": "1.2.2",
                            "name": "Community visits",
                            "disaggregation": [
                                {"name": "Adults", "target": 200, "months": [20, 20, 20, 10, 10, 10, 5, 5, 5, 0, 0, 0]},
                                {"name": "Children", "target": 150, "months": [15, 15, 15, 5, 5, 5, 5, 5, 5, 0, 0, 0]},
                                {"name": "Seniors", "target": 50, "months": [5, 5, 5, 5, 5, 5, 5, 5, 5, 0, 0, 0]},
                            ]
                        }
                    ]
                },
                {
                    "name": "Output 3: Remote counseling",
                    "indicators": [
                        {
                            "code": "1.3.1",
                            "name": "Online sessions",
                            "disaggregation": [
                                {"name": "Male", "target": 60, "months": [10, 10, 10, 10, 10, 10, 5, 5, 0, 0, 0, 0]},
                                {"name": "Female", "target": 80, "months": [15, 15, 15, 10, 10, 15, 5, 5, 0, 0, 0, 0]},
                                {"name": "Teenagers", "target": 40, "months": [5, 5, 5, 5, 5, 5, 5, 5, 0, 0, 0, 0]},
                            ]
                        }
                    ]
                }
            ]
        },
        {
            "name": "Communities have improved mental health knowledge",
            "outputs": [
                {
                    "name": "Output 4: Training & workshops",
                    "indicators": [
                        {
                            "code": "2.1.1",
                            "name": "Mental health workshops",
                            "disaggregation": [
                                {"name": "Male participants", "target": 60, "months": [10, 10, 10, 10, 10, 10, 5, 5, 0, 0, 0, 0]},
                                {"name": "Female participants", "target": 80, "months": [15, 15, 15, 10, 10, 15, 5, 5, 0, 0, 0, 0]},
                                {"name": "Youth", "target": 50, "months": [5, 5, 5, 5, 5, 5, 5, 5, 0, 0, 0, 0]},
                            ]
                        },
                        {
                            "code": "2.1.2",
                            "name": "Workshops follow-up",
                            "disaggregation": [
                                {"name": "Adults", "target": 100, "months": [10, 10, 10, 10, 10, 10, 10, 10, 10, 0, 0, 0]},
                                {"name": "Children", "target": 120, "months": [12, 12, 12, 12, 12, 12, 12, 12, 12, 0, 0, 0]},
                                {"name": "Seniors", "target": 30, "months": [3, 3, 3, 3, 3, 3, 3, 3, 3, 0, 0, 0]},
                            ]
                        }
                    ]
                }
            ]
        },
        {
            "name": "Communities have improved psychosocial support",
            "outputs": [
                {
                    "name": "Output 5: Group therapy",
                    "indicators": [
                        {
                            "code": "3.1.1",
                            "name": "Therapy sessions",
                            "disaggregation": [
                                {"name": "Men", "target": 60, "months": [10, 10, 10, 5, 5, 5, 5, 5, 5, 0, 0, 0]},
                                {"name": "Women", "target": 80, "months": [15, 15, 15, 10, 10, 15, 5, 5, 0, 0, 0, 0]},
                                {"name": "Teenagers", "target": 50, "months": [5, 5, 5, 5, 5, 5, 5, 5, 0, 0, 0, 0]},
                            ]
                        }
                    ]
                },
                {
                    "name": "Output 6: Counseling calls",
                    "indicators": [
                        {
                            "code": "3.2.1",
                            "name": "Hotline calls",
                            "disaggregation": [
                                {"name": "Male callers", "target": 100, "months": [10, 10, 10, 10, 10, 10, 10, 10, 10, 0, 0, 0]},
                                {"name": "Female callers", "target": 150, "months": [15, 15, 15, 15, 15, 15, 15, 15, 15, 0, 0, 0]},
                                {"name": "Children callers", "target": 50, "months": [5, 5, 5, 5, 5, 5, 5, 5, 5, 0, 0, 0]},
                            ]
                        }
                    ]
                }
            ]
        }
    ]
}



# -------------------------
# ساخت شیت اکسل
# -------------------------
wb = Workbook()
ws = wb.active
ws.title = "Monitoring"

headers = [
    "Impact/Goal (LFA)", "Result/Outcome", "Outputs", "Indicators",
    "Target", "Total achievement", "% Achieved",
    "Q1", "Jan", "Feb", "Mar",
    "Q2", "Apr", "May", "Jun",
    "Q3", "Jul", "Aug", "Sep",
    "Q4", "Oct", "Nov", "Dec"
]
ws.append(headers)

# رنگ آبی آسمانی برای ردیف indicator
indicator_fill = PatternFill(start_color="87CEEB", end_color="87CEEB", fill_type="solid")

# header style
header_fill = PatternFill(start_color="317C43", end_color="317C43", fill_type="solid")
for col in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(col)].width = 12

thin = Side(border_style="thin", color="CCCCCC")

# helper: ensure months array length 12
def ensure_12(months):
    if not months:
        return [0] * 12
    a = list(months)[:12]
    if len(a) < 12:
        a += [0] * (12 - len(a))
    return [int(x or 0) for x in a]

def quarter_sums_from_list(months12):
    return [sum(months12[0:3]), sum(months12[3:6]), sum(months12[6:9]), sum(months12[9:12])]

row = 2
impact_start_row = row
impact_written = False

quarter_fill = PatternFill(start_color="87CEEB", end_color="87CEEB", fill_type="solid")


for outcome in final_apr["outcomes"]:
    outcome_start_row = row
    outcome_written = False

    for output in outcome["outputs"]:
        output_start_row = row
        output_written = False

        for indicator in output["indicators"]:
            indicator_start_row = row

            # ---------- جمع disaggregation ----------
            months_sum = [0] * 12
            total_target = 0
            for d in indicator["disaggregation"]:
                dm = ensure_12(d.get("months", []))
                months_sum = [months_sum[i] + dm[i] for i in range(12)]
                total_target += int(d.get("target", 0) or 0)

            quarters = quarter_sums_from_list(months_sum)
            total_achieved = sum(months_sum)
            percent_fraction = (total_achieved / total_target) if total_target else 0.0

            # ---------- نوشتن ردیف indicator ----------
            indicator_label = f"{indicator.get('code','')} - {indicator.get('name','')}"
            ws.append([
                final_apr["impact"] if not impact_written else "",
                outcome["name"] if not outcome_written else "",
                output["name"] if not output_written else "",
                indicator_label,
                total_target,
                total_achieved,
                percent_fraction,
                quarters[0], months_sum[0], months_sum[1], months_sum[2],
                quarters[1], months_sum[3], months_sum[4], months_sum[5],
                quarters[2], months_sum[6], months_sum[7], months_sum[8],
                quarters[3], months_sum[9], months_sum[10], months_sum[11]
            ])
            ws.cell(row=row, column=7).number_format = '0.00%'

            # ---------- رنگ‌آمیزی indicator row و ستون Q1-Q4 ----------
            for col in list(range(4, 8)) + list(range(8, 24)):
                ws.cell(row=indicator_start_row, column=col).fill = indicator_fill

            row += 1

            # ---------- نوشتن disaggregation ها ----------
            for d in indicator["disaggregation"]:
                dm = ensure_12(d.get("months", []))
                d_quarters = quarter_sums_from_list(dm)
                d_total = sum(dm)
                d_target = int(d.get("target", 0) or 0)
                d_percent_fraction = (d_total / d_target) if d_target else 0.0

                ws.append([
                    "", "", "", f"   {d['name']}",
                    d_target,
                    d_total,
                    d_percent_fraction,
                    d_quarters[0], dm[0], dm[1], dm[2],
                    d_quarters[1], dm[3], dm[4], dm[5],
                    d_quarters[2], dm[6], dm[7], dm[8],
                    d_quarters[3], dm[9], dm[10], dm[11]
                ])
                ws.cell(row=row, column=7).number_format = '0.00%'


                for col in [8, 12, 16, 20]:
                    ws.cell(row=row, column=col).fill = quarter_fill

                row += 1

            # mark that we've written outcome/output/impact at least once
            impact_written = True
            outcome_written = True
            output_written = True

        # merge output column اگر چند ردیف داشت
        if row - output_start_row > 1:
            ws.merge_cells(start_row=output_start_row, end_row=row - 1, start_column=3, end_column=3)

    # merge outcome column اگر چند ردیف داشت
    if row - outcome_start_row > 1:
        ws.merge_cells(start_row=outcome_start_row, end_row=row - 1, start_column=2, end_column=2)

# merge impact column اگر چند ردیف داشت
if row - impact_start_row > 1:
    ws.merge_cells(start_row=impact_start_row, end_row=row - 1, start_column=1, end_column=1)


# ---------- فرمت‌بندی نهایی ----------
thin_border = Border(left=Side(style='thin', color='CCCCCC'),
                     right=Side(style='thin', color='CCCCCC'),
                     top=Side(style='thin', color='CCCCCC'),
                     bottom=Side(style='thin', color='CCCCCC'))

numeric_cols = list(range(5, 5 + 1))  # Target col=5
numeric_cols += [6]  # Total achievement
numeric_cols += [7]  # % Achieved
numeric_cols += list(range(8, 24))  # quarters + months

for r in ws.iter_rows(min_row=1, max_row=row - 1, min_col=1, max_col=23):
    for cell in r:
        cell.border = thin_border
        if cell.column in numeric_cols:
            cell.alignment = Alignment(horizontal='right', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

# set column widths
widths = [30, 30, 35, 45, 12, 15, 12] + [9]*16
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

# save
wb.save("Monitoring_Final_fixed.xlsx")
print("✅ Monitoring_Final_fixed.xlsx ساخته شد — indicator row آبی شد، merge و محاسبات درست است.")
