import json
from http.server import BaseHTTPRequestHandler, HTTPServer
from io import BytesIO
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HOST = "0.0.0.0"
PORT = 5000


def ensure_months(months):
    if not months:
        return []
    return [int(value or 0) for value in months]


def quarter_sum(months, start):
    return sum(months[start : start + 3])


def build_workbook(payload):
    wb = Workbook()
    ws = wb.active
    ws.title = "Monitoring"

    headers = [
        "Impact/Goal (LFA)",
        "Result/Outcome",
        "Outputs",
        "Indicators",
        "Target",
        "Total achievement",
        "% Achieved",
    ]

    max_months = 0
    for outcome in payload.get("outcomes", []):
        for output in outcome.get("outputs", []):
            for indicator in output.get("indicators", []):
                for disaggregation in indicator.get("disaggregation", []):
                    max_months = max(max_months, len(disaggregation.get("months", []) or []))

    for month_index in range(max_months):
        if month_index % 3 == 0:
            headers.append(f"Q{(month_index // 3) + 1}")
        headers.append(f"M{month_index + 1}")

    ws.append(headers)

    header_fill = PatternFill(start_color="317C43", end_color="317C43", fill_type="solid")
    indicator_fill = PatternFill(start_color="87CEEB", end_color="87CEEB", fill_type="solid")
    border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = 14

    impact = payload.get("impact") or payload.get("projectGoal") or ""
    row = 2
    impact_written = False

    for outcome in payload.get("outcomes", []):
        outcome_written = False
        for output in outcome.get("outputs", []):
            output_written = False
            for indicator in output.get("indicators", []):
                disaggregations = indicator.get("disaggregation", []) or []
                months_sum = [0] * max_months
                total_target = 0

                for disaggregation in disaggregations:
                    months = ensure_months(disaggregation.get("months", []))
                    total_target += int(disaggregation.get("target", 0) or 0)
                    for index, value in enumerate(months):
                        months_sum[index] += value

                total_achievement = sum(months_sum)
                percent = (total_achievement / total_target) if total_target else 0
                indicator_label = f"{indicator.get('code', '')} - {indicator.get('name', '')}".strip(" -")

                values = [
                    impact if not impact_written else "",
                    outcome.get("name", "") if not outcome_written else "",
                    output.get("name", "") if not output_written else "",
                    indicator_label,
                    total_target,
                    total_achievement,
                    percent,
                ]

                for month_index, value in enumerate(months_sum):
                    if month_index % 3 == 0:
                        values.append(quarter_sum(months_sum, month_index))
                    values.append(value)

                ws.append(values)
                ws.cell(row=row, column=7).number_format = "0.00%"
                for col in range(4, len(headers) + 1):
                    ws.cell(row=row, column=col).fill = indicator_fill
                row += 1

                for disaggregation in disaggregations:
                    months = ensure_months(disaggregation.get("months", []))
                    padded_months = months + [0] * max(0, max_months - len(months))
                    target = int(disaggregation.get("target", 0) or 0)
                    total = sum(padded_months)
                    dis_percent = (total / target) if target else 0

                    values = ["", "", "", f"   {disaggregation.get('name', '')}", target, total, dis_percent]
                    for month_index, value in enumerate(padded_months):
                        if month_index % 3 == 0:
                            values.append(quarter_sum(padded_months, month_index))
                        values.append(value)

                    ws.append(values)
                    ws.cell(row=row, column=7).number_format = "0.00%"
                    row += 1

                impact_written = True
                outcome_written = True
                output_written = True

    for excel_row in ws.iter_rows(min_row=1, max_row=max(1, row - 1), min_col=1, max_col=len(headers)):
        for cell in excel_row:
            cell.border = border
            cell.alignment = Alignment(
                horizontal="right" if cell.column >= 5 else "left",
                vertical="center",
                wrap_text=True,
            )

    widths = [30, 30, 35, 45, 12, 18, 12] + [10] * max(0, len(headers) - 7)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


class ExcelHandler(BaseHTTPRequestHandler):
    def _set_cors_headers(self):
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type, Authorization")

    def do_OPTIONS(self):
        self.send_response(204)
        self._set_cors_headers()
        self.end_headers()

    def do_GET(self):
        if urlparse(self.path).path != "/health":
            self.send_response(404)
            self.end_headers()
            return

        self.send_response(200)
        self._set_cors_headers()
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(json.dumps({"status": True}).encode("utf-8"))

    def do_POST(self):
        if urlparse(self.path).path != "/generate-excel":
            self.send_response(404)
            self.end_headers()
            return

        try:
            content_length = int(self.headers.get("Content-Length", 0))
            payload = json.loads(self.rfile.read(content_length) or b"{}")
            workbook = build_workbook(payload)

            self.send_response(200)
            self._set_cors_headers()
            self.send_header(
                "Content-Type",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            self.send_header("Content-Disposition", 'attachment; filename="Monitoring_Dynamic.xlsx"')
            self.end_headers()
            self.wfile.write(workbook)
        except Exception as exc:
            self.send_response(500)
            self._set_cors_headers()
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(json.dumps({"status": False, "message": str(exc)}).encode("utf-8"))


if __name__ == "__main__":
    HTTPServer((HOST, PORT), ExcelHandler).serve_forever()
